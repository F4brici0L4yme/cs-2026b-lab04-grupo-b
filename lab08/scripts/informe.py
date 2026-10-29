"""Maqueta el informe con el estilo compacto de los informes 06 y 07 del grupo.

El Markdown sigue siendo la fuente editable. Se reutiliza el encabezado y la
tabla institucional de la plantilla, sin alterar el documento original.
"""

import argparse
from copy import deepcopy
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import html
import markdown
from PIL import Image

BASE = Path(__file__).resolve().parents[1]
FUENTE = BASE / "docs/informe.md"
DESTINO = BASE / "informe/Informe_Lab08_GrupoB.docx"


def texto(elemento):
    return "".join(elemento.itertext()).strip()


def inline(parrafo, elemento):
    if elemento.text:
        parrafo.add_run(re.sub(r"\s+", " ", elemento.text))
    for hijo in elemento:
        run = parrafo.add_run(re.sub(r"\s+", " ", texto(hijo)))
        if hijo.tag in ("strong", "b"):
            run.bold = True
        elif hijo.tag in ("em", "i"):
            run.italic = True
        elif hijo.tag == "code":
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        if hijo.tail:
            parrafo.add_run(re.sub(r"\s+", " ", hijo.tail))


def sombrear(celda, color):
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:fill"), color)
    celda._tc.get_or_add_tcPr().append(sombra)


def espaciar_celda(celda, arriba=90, abajo=90, lateral=105):
    """Márgenes en twips: evita texto pegado a bordes aun en filas altas."""
    celda.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    propiedades = celda._tc.get_or_add_tcPr()
    margenes = propiedades.first_child_found_in("w:tcMar")
    if margenes is None:
        margenes = OxmlElement("w:tcMar")
        propiedades.append(margenes)
    for lado, valor in (("top", arriba), ("bottom", abajo),
                        ("left", lateral), ("right", lateral)):
        nodo = margenes.find(qn(f"w:{lado}"))
        if nodo is None:
            nodo = OxmlElement(f"w:{lado}")
            margenes.append(nodo)
        nodo.set(qn("w:w"), str(valor))
        nodo.set(qn("w:type"), "dxa")


def separar_tabla(tabla):
    """Las tablas anidadas necesitan sangría propia: ignoran el margen del marco."""
    tabla.autofit = False
    propiedades = tabla._tbl.tblPr
    sangria = OxmlElement("w:tblInd")
    sangria.set(qn("w:w"), "283")
    sangria.set(qn("w:type"), "dxa")
    propiedades.append(sangria)
    for columna in tabla.columns:
        columna.width = Cm(16.1 / len(tabla.columns))


SECCIONES = ("SOLUCIÓN Y RESULTADOS", "SOLUCIÓN DEL CUESTIONARIO",
             "CONCLUSIONES", "REFERENCIAS Y BIBLIOGRAFÍA")


def enmarcar_secciones(doc):
    """Crea un marco independiente por sección, sin tocar la portada."""
    cuerpo = doc._element.body
    inicios = {}
    for elemento in cuerpo:
        if elemento.tag == qn("w:tbl"):
            contenido = "".join(elemento.itertext())
            for titulo in SECCIONES:
                if titulo in contenido:
                    inicios[titulo] = elemento
    if len(inicios) != len(SECCIONES):
        raise ValueError("Falta una barra de sección en el informe")

    for i, titulo in enumerate(SECCIONES):
        inicio = inicios[titulo]
        fin = inicios[SECCIONES[i + 1]] if i + 1 < len(SECCIONES) else cuerpo.sectPr
        fragmentos = []
        actual = inicio
        while actual is not fin:
            fragmentos.append(actual)
            actual = actual.getnext()
        marco = doc.add_table(rows=2, cols=1)
        marco.alignment = WD_TABLE_ALIGNMENT.CENTER
        cuerpo.insert(list(cuerpo).index(inicio), marco._element)
        cabecera, celda = marco.cell(0, 0), marco.cell(1, 0)
        sombrear(cabecera, "C73514")
        espaciar_celda(cabecera, arriba=45, abajo=45, lateral=0)
        # La barra ocupa toda la fila superior; el margen lateral solo se
        # aplica a la fila del contenido, no al título.
        cabecera._tc.replace(cabecera.paragraphs[0]._p,
                            deepcopy(inicio.find(".//" + qn("w:p"))))
        cuerpo.remove(inicio)
        marco.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for fila, lados in ((cabecera, ("top", "left", "right")),
                            (celda, ("bottom", "left", "right"))):
            bordes = OxmlElement("w:tcBorders")
            for lado in lados:
                borde = OxmlElement(f"w:{lado}")
                borde.set(qn("w:val"), "single")
                borde.set(qn("w:sz"), "6")
                borde.set(qn("w:color"), "222222")
                bordes.append(borde)
            fila._tc.get_or_add_tcPr().append(bordes)
        # 0,5 cm entre el borde y el contenido, también en cada nueva sección.
        espaciar_celda(celda, arriba=60, abajo=80, lateral=283)
        for elemento in fragmentos[1:]:
            celda._tc.insert(-1, elemento)
        if i + 1 < len(SECCIONES):
            espacio = doc.add_paragraph()
            if SECCIONES[i + 1] in ("SOLUCIÓN DEL CUESTIONARIO", "CONCLUSIONES"):
                # Evitar barras huérfanas y que las conclusiones se partan.
                espacio.paragraph_format.page_break_before = True
                espacio.paragraph_format.space_after = Pt(0)
                espacio.add_run(" ").font.size = Pt(1)
            else:
                espacio.paragraph_format.space_after = Pt(7)
                espacio.paragraph_format.line_spacing = 0.5
            cuerpo.insert(list(cuerpo).index(fin), espacio._p)


def barra(doc, titulo):
    tabla = doc.add_table(rows=1, cols=1)
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    separar_tabla(tabla)
    celda = tabla.cell(0, 0)
    espaciar_celda(celda, arriba=45, abajo=45)
    sombrear(celda, "C73514")
    parrafo = celda.paragraphs[0]
    parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    parrafo.paragraph_format.space_after = Pt(0)
    parrafo.paragraph_format.space_before = Pt(0)
    parrafo.paragraph_format.keep_with_next = True
    run = parrafo.add_run(titulo)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(255, 255, 255)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def configurar(doc):
    normal = next(s for s in doc.styles if s.style_id == "Normal")
    normal.font.name = "Arial"
    normal.font.size = Pt(9)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1.04
    for nombre in ("Heading 1", "Heading 2", "Heading 3"):
        estilo = doc.styles[nombre]
        estilo.font.name = "Arial"
        estilo.font.size = Pt(10 if nombre == "Heading 3" else 11)
        estilo.font.bold = True
        estilo.font.color.rgb = RGBColor.from_string("C73514")
        estilo.paragraph_format.space_before = Pt(6)
        estilo.paragraph_format.space_after = Pt(3)
        estilo.paragraph_format.keep_with_next = True
    if not any(s.name == "Table Grid" for s in doc.styles):
        doc.styles.element.append(deepcopy(Document().styles["Table Grid"].element))
    for seccion in doc.sections:
        seccion.page_width, seccion.page_height = Cm(21), Cm(29.7)
        seccion.left_margin = seccion.right_margin = Cm(1.85)
        seccion.top_margin, seccion.bottom_margin = Cm(3.2), Cm(1.25)
    # La plantilla ya numera las páginas en el encabezado institucional.


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--plantilla", type=Path)
    args = parser.parse_args()
    doc = Document(args.plantilla) if args.plantilla and args.plantilla.exists() else Document()
    if args.plantilla and args.plantilla.exists():
        tabla = doc.tables[0]
        tabla.cell(1, 1).text = "Construcción de Software"
        tabla.cell(2, 1).text = "Diseñar la arquitectura guiada por atributos de calidad: ADD, ATAM-lite y fitness functions — RutaSIT Arequipa"
        tabla.cell(3, 1).text = "08"
        tabla.cell(3, 3).text = "2026-B"
        tabla.cell(3, 5).text = "VI"
        tabla.cell(4, 1).text = "31/10/2026 (fecha límite)"
        tabla.cell(4, 3).text = "23:59"
        tabla.cell(5, 0).text = ("INTEGRANTES:\nLayme Salas, Rodrigo Fabricio\n"
                                "Auccacusi Conde, Brayan Carlos\nCastillo Lazo, Rodrigo Zarun")
        tabla.cell(5, 5).text = "Reservado al docente"
        tabla.cell(6, 0).text = "DOCENTE:\nMg. Antonio Arroyo Paz"
        for fila in tabla.rows:
            for celda in fila.cells:
                for parrafo in celda.paragraphs:
                    for run in parrafo.runs:
                        run.font.name, run.font.size = "Arial", Pt(9)
        # Eliminar cuerpos vacíos de solución/retroalimentación/referencias,
        # conservando portada, relaciones de imágenes y encabezados de plantilla.
        encontrado = False
        for hijo in list(doc._element.body):
            if hijo is tabla._element:
                encontrado = True
            elif encontrado and hijo.tag != qn("w:sectPr"):
                doc._element.body.remove(hijo)
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
    else:
        doc.add_heading("UNIVERSIDAD NACIONAL DE SAN AGUSTÍN", 0)
        doc.add_paragraph("Escuela Profesional de Ingeniería de Sistemas\nINFORME DE LABORATORIO 08")
    configurar(doc)
    arbol = html.fragment_fromstring(markdown.markdown(FUENTE.read_text(),
                                                       extensions=["tables", "fenced_code"]),
                                      create_parent="div")
    introduccion = True
    for nodo in arbol:
        if introduccion:
            if nodo.tag == "h2" and texto(nodo) == "Resumen":
                introduccion = False
                barra(doc, "SOLUCIÓN Y RESULTADOS")
                doc.add_heading("RESUMEN", level=2)
            continue
        if nodo.tag in ("h1", "h2", "h3", "h4"):
            titulo = texto(nodo)
            if nodo.tag == "h2" and titulo.startswith(("I. ", "II. ", "III. ", "IV. ")):
                if titulo.startswith("I. "):
                    continue  # el desarrollo comenzó después de la portada
                if titulo.startswith("II. "):
                    barra(doc, "SOLUCIÓN DEL CUESTIONARIO")
                elif titulo.startswith("III. "):
                    barra(doc, "CONCLUSIONES")
                else:
                    barra(doc, "REFERENCIAS Y BIBLIOGRAFÍA")
            else:
                doc.add_heading(titulo, level=min(int(nodo.tag[1]), 3))
        elif nodo.tag == "table":
            filas = nodo.xpath(".//tr")
            tabla = doc.add_table(rows=0, cols=len(filas[0]))
            tabla.style = "Table Grid"
            tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
            separar_tabla(tabla)
            for i, fila in enumerate(filas):
                celdas = tabla.add_row().cells
                tabla.rows[-1]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
                for celda, valor in zip(celdas, fila):
                    espaciar_celda(celda)
                    inline(celda.paragraphs[0], valor)
                    for parrafo in celda.paragraphs:
                        for run in parrafo.runs:
                            run.font.size = Pt(8)
                            if i == 0:
                                run.bold = True
                    if i == 0:
                        sombrear(celda, "C73514")
                        for run in celda.paragraphs[0].runs:
                            run.font.color.rgb = RGBColor(255, 255, 255)
                    elif i % 2 == 0:
                        sombrear(celda, "F1F1F1")
                if i == 0:
                    repetir = OxmlElement("w:tblHeader")
                    tabla.rows[0]._tr.get_or_add_trPr().append(repetir)
            doc.add_paragraph().paragraph_format.space_after = Pt(0)
        elif nodo.tag == "p" and nodo.find("img") is not None:
            imagen = nodo.find("img")
            ruta = (FUENTE.parent / imagen.get("src")).resolve()
            if not ruta.exists():
                raise FileNotFoundError(ruta)
            with Image.open(ruta) as im:
                ancho, alto = im.size
            # Ajustar también por altura para no desbordar una página.
            ancho_cm = min(15.8, 13.0 * ancho / alto)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(ruta), width=Cm(ancho_cm))
            cap = doc.add_paragraph(imagen.get("alt", "Figura"))
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in cap.runs:
                run.italic = True
                run.font.size = Pt(8)
        elif nodo.tag in ("ul", "ol"):
            for n, item in enumerate(nodo, 1):
                p = doc.add_paragraph()
                p.add_run(f"{n}. " if nodo.tag == "ol" else "• ")
                inline(p, item)
        elif nodo.tag == "pre":
            p = doc.add_paragraph()
            run = p.add_run(texto(nodo))
            run.font.name, run.font.size = "Consolas", Pt(8)
        else:
            inline(doc.add_paragraph(), nodo)
    if args.plantilla and args.plantilla.exists():
        enmarcar_secciones(doc)
    DESTINO.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DESTINO)
    print(DESTINO)


if __name__ == "__main__":
    main()
